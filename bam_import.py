"""
bam_import.py — Import der amtlichen Gefahrgutdaten der BAM (Datenbank GEFAHRGUT)

Warum dieses Modul existiert
───────────────────────────
Die ursprüngliche Datenbasis wurde per Heuristik aus dem ADR-PDF geparst
(„nächstgelegene Ziffer 0–4" als Beförderungskategorie, Default Kategorie 3
wenn nichts gefunden wurde). Das erzeugt stille Fehler in einer Anwendung,
die über eine Freistellung nach ADR 1.1.3.6 entscheidet.

Die BAM (Bundesanstalt für Materialforschung und -prüfung) veröffentlicht
dieselben Daten als strukturierte Datei — redaktionell gepflegt, vollständig
und seit 2025-07-23 kostenfrei. Damit entfällt das Raten vollständig:
Beförderungskategorie, Tunnelcode und sogar der Punktfaktor (1.1.3.6) stehen
als eigene Felder bereit.

Datenquelle
  https://tes.bam.de/en/dangerous-goods-database/products/dangerous-goods-dataservice
  Datensatz „UN-No. System / ADR"  →  dgg-daten-adr-un.zip  →  ADR25.xlsx / ADR25_csv.txt

Lizenz (pflichtangaben!)
  Datenlizenz Deutschland – Namensnennung – Version 2.0  (dl-de/by-2-0)
  Quellenangabe:
    „Source: Bundesanstalt für Materialforschung und -prüfung (BAM) –
     Datenbank GEFAHRGUT – URL: tes.bam.de/TES/Navigation/EN/DGG-Database/
     dgg-database.html — Data licence Germany – attribution – Version 2.0"
  Der Aufruf von `attribution()` liefert diesen Text fertig formatiert.

  ACHTUNG: Die BAM untersagt gemäss § 44b (3) UrhG die Nutzung dieser Daten
  für Text- und Data-Mining. Die Verwendung als Nachschlagetabelle in dieser
  Anwendung ist eine normale Datenweiternutzung und zulässig; ein Training
  von Modellen mit diesen Daten benötigt die schriftliche Zustimmung der BAM.

Wichtiges Datenmodell-Wissen
───────────────────────────
Eine UN-Nummer hat in Tabelle A **mehrere Varianten** (verschiedene
Verpackungsgruppen, Spezifikationen, Klassifizierungscodes). Die
Beförderungskategorie hängt von der Variante ab:

    UN 1133 KLEBSTOFFE   PG I → Kat 1 | PG II → Kat 2 | PG III → Kat 3

Deshalb ist der natürliche Schlüssel (un_number, variant) — nicht un_number
allein. Ein UPDATE ... WHERE un_number = ? würde die Kategorien der
Verpackungsgruppen gegenseitig überschreiben und falsche Punktzahlen
erzeugen.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, List, Optional

# Spaltennamen der BAM-Datei (ADR25.xlsx / ADR25_csv.txt), 97 Felder.
COL = {
    "un": "S_UNNR",
    "variant": "N_LFDNR",
    "prefix": "S_VORSILBE",
    "name": "S_NAME",
    "spec": "S_SPEZIFIKATION",
    "klasse": "S_KLASSE",
    "klass_code": "S_KLASSIFIZIERUNGSCODE",
    "pg": "S_VP_GRUPPE",
    "kenn": ["S_KENN1", "S_KENN2", "S_KENN3", "S_KENN4"],
    "sv": ["S_SV1", "S_SV2", "S_SV3", "S_SV4", "S_SV5", "S_SV6",
           "S_SV7", "S_SV8", "S_SV9", "SV_10", "SV_11"],
    "lq": "S_BEGRENZTE_MENGEN",
    "eq": "S_FREIGEST_MENGEN",
    "verpack": [f"S_VERPACKUNGSANW{i}" for i in range(1, 10)],
    "tank": "S_TANK_CODE1",
    "tankfahrzeug": "S_TANKFAHRZEUG",
    "kategorie": "S_BEFOERDERUNGSKATEGORIE",
    "tunnel": "S_TUNNEL_CODE",
    "gefahrnr": "S_GEFAHRNR",
    "multiplikator": "N_MULTIPLIKATOR",
    "verbot": "N_VERBOT",
    "name_en": "S_NAME_E",
    "spec_en": "S_SPEZIFIKATION_E",
    "name_fr": "S_NAME_FR",
    "spec_fr": "S_SPEZIFIKATION_FR",
    "bemerkung": "S_BEMERKUNG",
}

# Die BAM-Datei ist cp1252-kodiert (kein UTF-8!). xlsx liefert ohnehin str.
ENCODINGS = ("cp1252", "latin-1", "utf-8")

# Platzhalter, die die BAM für „nicht vorhanden" verwendet.
NULL_TOKENS = frozenset({"", "-"})


def attribution() -> str:
    """Vorgeschriebene Quellenangabe für die Weiternutzung der BAM-Daten.

    Muss in der Anwendung sichtbar ausgegeben werden (Impressum / „Datenquellen").
    """
    return (
        "Source: Bundesanstalt für Materialforschung und -prüfung (BAM) – "
        "Datenbank GEFAHRGUT – URL: tes.bam.de/TES/Navigation/EN/DGG-Database/"
        "dgg-database.html — Data licence Germany – attribution – Version 2.0"
    )


def _clean(value: Any) -> Optional[str]:
    """Normalisiert einen Wert: None bei leer oder BAM-Platzhalter '-'."""
    if value is None:
        return None
    if isinstance(value, float):
        # openpyxl liefert ganzzahlige Floats (z. B. 50.0 für den Faktor)
        if value.is_integer():
            value = int(value)
        else:
            value = repr(value)
    text = str(value).replace("\xa0", " ").strip()
    if text in NULL_TOKENS:
        return None
    return text


def _join(row: Dict[str, Any], keys: Iterable[str], sep: str = ",") -> Optional[str]:
    parts = [_clean(row.get(k)) for k in keys]
    parts = [p for p in parts if p]
    if not parts:
        return None
    # Doppelte erhalten die Reihenfolge (dict.fromkeys)
    return sep.join(dict.fromkeys(parts))


@dataclass
class BamEntry:
    """Ein Eintrag (eine Variante) aus der BAM-UN-Nummern-Liste."""

    un_number: str
    variant: Optional[int] = None
    name_de: str = ""
    spec_de: Optional[str] = None
    name_en: Optional[str] = None
    name_fr: Optional[str] = None
    hazard_class: Optional[str] = None
    classification_code: Optional[str] = None
    packing_group: Optional[str] = None
    labels: Optional[str] = None            # Gefahrzettel (5)
    special_provisions: Optional[str] = None  # Sondervorschriften (6)
    limited_quantity: Optional[str] = None   # (7a) LQ
    excepted_quantity: Optional[str] = None  # (7b) EQ
    packing_instructions: Optional[str] = None  # (8)
    tank_code: Optional[str] = None          # (12)
    vehicle_for_tank: Optional[str] = None   # (14)
    transport_category: Optional[int] = None  # (15) — der Schlüsselwert
    tunnel_code: Optional[str] = None        # (15)
    hazard_identification_no: Optional[str] = None  # (20) Kemler-Zahl
    multiplier: Optional[float] = None       # Faktor nach 1.1.3.6 (BAM-berechnet)
    prohibited: Optional[str] = None
    notes: Optional[str] = None

    @property
    def full_name_de(self) -> str:
        if self.spec_de:
            return f"{self.name_de}, {self.spec_de}"
        return self.name_de


def _row_to_entry(row: Dict[str, Any]) -> Optional[BamEntry]:
    un = _clean(row.get(COL["un"]))
    if not un:
        return None
    un = un.strip().zfill(4) if un.isdigit() else un.strip()

    name = _clean(row.get(COL["name"])) or ""
    prefix = _clean(row.get(COL["prefix"]))
    if prefix:
        name = f"{prefix} {name}".strip()

    variant_raw = _clean(row.get(COL["variant"]))
    variant = None
    if variant_raw and variant_raw.isdigit():
        variant = int(variant_raw)

    cat_raw = _clean(row.get(COL["kategorie"]))
    category = None
    if cat_raw and cat_raw.isdigit() and 0 <= int(cat_raw) <= 4:
        category = int(cat_raw)

    mult_raw = _clean(row.get(COL["multiplikator"]))
    multiplier = None
    if mult_raw:
        try:
            multiplier = float(mult_raw.replace(",", "."))
        except ValueError:
            multiplier = None

    tunnel = _clean(row.get(COL["tunnel"]))
    if tunnel:
        tunnel = tunnel.strip().strip("()").strip()
        if tunnel in NULL_TOKENS:
            tunnel = None

    name_en = _clean(row.get(COL["name_en"]))
    spec_en = _clean(row.get(COL["spec_en"]))
    if name_en and spec_en:
        name_en = f"{name_en}, {spec_en}"

    name_fr = _clean(row.get(COL["name_fr"]))
    spec_fr = _clean(row.get(COL["spec_fr"]))
    if name_fr and spec_fr:
        name_fr = f"{name_fr}, {spec_fr}"

    return BamEntry(
        un_number=un,
        variant=variant,
        name_de=name,
        spec_de=_clean(row.get(COL["spec"])),
        name_en=name_en,
        name_fr=name_fr,
        hazard_class=_clean(row.get(COL["klasse"])),
        classification_code=_clean(row.get(COL["klass_code"])),
        packing_group=_clean(row.get(COL["pg"])),
        labels=_join(row, COL["kenn"]),
        special_provisions=_join(row, COL["sv"]),
        limited_quantity=_clean(row.get(COL["lq"])),
        excepted_quantity=_clean(row.get(COL["eq"])),
        packing_instructions=_join(row, COL["verpack"]),
        tank_code=_clean(row.get(COL["tank"])),
        vehicle_for_tank=_clean(row.get(COL["tankfahrzeug"])),
        transport_category=category,
        tunnel_code=tunnel,
        hazard_identification_no=_clean(row.get(COL["gefahrnr"])),
        multiplier=multiplier,
        prohibited=_clean(row.get(COL["verbot"])),
        notes=_clean(row.get(COL["bemerkung"])),
    )


# ── Lesen der Quelldateien ────────────────────────────────────────────

def _read_xlsx(source) -> List[Dict[str, Any]]:
    """Liest das BAM-xlsx. `source` = Pfad oder FileStorage/BytesIO."""
    from openpyxl import load_workbook

    if hasattr(source, "read"):
        data = source.read()
        if isinstance(data, str):
            data = data.encode("utf-8")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    else:
        wb = load_workbook(source, read_only=True, data_only=True)

    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header: Optional[List[str]] = None
        out: List[Dict[str, Any]] = []
        for raw in rows:
            if raw is None:
                continue
            cells = [("" if c is None else c) for c in raw]
            if header is None:
                # Kopfzeile: erste Zeile, die den UN-Feldnamen enthält
                if any(str(c).strip() == COL["un"] for c in cells):
                    header = [str(c).strip() for c in cells]
                continue
            if all(str(c).strip() == "" for c in cells):
                continue
            out.append({h: cells[i] if i < len(cells) else None
                        for i, h in enumerate(header)})
        if header is None:
            raise ValueError("Kopfzeile mit Spalte S_UNNR nicht gefunden")
        return out
    finally:
        wb.close()


def _read_csv(source) -> List[Dict[str, Any]]:
    """Liest die BAM-CSV (tabsepariert, cp1252). `source` = Pfad oder FileStorage."""
    if hasattr(source, "read"):
        blob = source.read()
        if isinstance(blob, str):
            blob = blob.encode("utf-8")
    else:
        with open(source, "rb") as fh:
            blob = fh.read()

    text = None
    for enc in ENCODINGS:
        try:
            text = blob.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = blob.decode("utf-8", errors="replace")

    lines = [ln for ln in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
             if ln.strip()]
    if not lines:
        raise ValueError("Leere CSV-Datei")

    sep = "\t" if "\t" in lines[0] else ";"
    header = [h.strip() for h in lines[0].split(sep)]
    if COL["un"] not in header:
        raise ValueError(
            f"Spalte {COL['un']} nicht gefunden — ist dies eine BAM-DGG-Datei? "
            f"Gefundene Spalten: {', '.join(header[:8])} ..."
        )

    out = []
    for ln in lines[1:]:
        cells = ln.split(sep)
        out.append({h: (cells[i].strip() if i < len(cells) else None)
                    for i, h in enumerate(header)})
    return out


def parse_bam_file(source, filename: str = "") -> List[BamEntry]:
    """Liest eine BAM-Datei (xlsx oder csv/txt) und liefert die Einträge.

    Args:
        source: Pfad, FileStorage oder BytesIO
        filename: Originaldateiname (nur für die Format-Erkennung nötig)

    Raises:
        ValueError: Wenn das Format nicht erkannt oder keine BAM-Datei ist.
    """
    name = (filename or getattr(source, "filename", "") or str(source)).lower()

    if name.endswith((".xlsx", ".xlsm")):
        rows = _read_xlsx(source)
    elif name.endswith((".csv", ".txt", ".tsv")):
        rows = _read_csv(source)
    else:
        # Unbekannte Endung: anhand der Magic Bytes entscheiden (ZIP = xlsx)
        if hasattr(source, "read"):
            pos = source.tell()
            head = source.read(4)
            source.seek(pos)
        else:
            with open(source, "rb") as fh:
                head = fh.read(4)
        rows = _read_xlsx(source) if head.startswith(b"PK\x03\x04") else _read_csv(source)

    entries: List[BamEntry] = []
    for row in rows:
        entry = _row_to_entry(row)
        if entry:
            entries.append(entry)
    return entries


# ── Plausibilitätsprüfung ─────────────────────────────────────────────

@dataclass
class BamCheck:
    """Ergebnis der Strukturprüfung einer BAM-Datei."""
    ok: bool = True
    entries: int = 0
    un_numbers: int = 0
    with_category: int = 0
    with_multiplier: int = 0
    problems: List[str] = field(default_factory=list)

    def as_dict(self) -> Dict[str, Any]:
        return {
            "ok": self.ok,
            "entries": self.entries,
            "un_numbers": self.un_numbers,
            "with_category": self.with_category,
            "with_multiplier": self.with_multiplier,
            "problems": self.problems,
        }


def check_entries(entries: List[BamEntry]) -> BamCheck:
    """Prüft die gelesenen Daten auf offensichtliche Strukturfehler.

    Erwartet werden für ADR 2025 rund 3.300 Varianten / ca. 2.350 UN-Nummern.
    Abweichungen sind ein Hinweis auf eine falsche oder abgeschnittene Datei,
    aber kein Ausschlussgrund — die Prüfung markiert nur.
    """
    chk = BamCheck(entries=len(entries))
    chk.un_numbers = len({e.un_number for e in entries})
    chk.with_category = sum(1 for e in entries if e.transport_category is not None)
    chk.with_multiplier = sum(1 for e in entries if e.multiplier is not None)

    if not entries:
        chk.ok = False
        chk.problems.append("Keine Einträge erkannt.")
        return chk

    if chk.un_numbers < 2000:
        chk.ok = False
        chk.problems.append(
            f"Nur {chk.un_numbers} UN-Nummern erkannt — für ADR werden "
            f"ca. 2.350 erwartet. Datei möglicherweise unvollständig."
        )

    cat_rate = chk.with_category / len(entries)
    if cat_rate < 0.9:
        chk.ok = False
        chk.problems.append(
            f"Nur {cat_rate:.0%} der Einträge haben eine Beförderungskategorie "
            f"(erwartet > 90 %)."
        )

    bad_un = [e.un_number for e in entries
              if not re.fullmatch(r"\d{4}", e.un_number or "")]
    if bad_un:
        chk.problems.append(
            f"{len(bad_un)} Einträge mit ungültiger UN-Nummer: "
            f"{', '.join(sorted(set(bad_un))[:5])}"
        )
        chk.ok = False

    dupes: Dict[str, int] = {}
    seen = set()
    for e in entries:
        key = (e.un_number, e.variant)
        if key in seen:
            dupes[f"{e.un_number}/{e.variant}"] = dupes.get(f"{e.un_number}/{e.variant}", 0) + 1
        seen.add(key)
    if dupes:
        chk.ok = False
        chk.problems.append(
            f"{len(dupes)} doppelte (UN-Nummer, Variante)-Paare, z. B. "
            f"{', '.join(list(dupes)[:5])}"
        )

    return chk
